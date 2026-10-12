"""Import a DealGate cost-margin workbook (.xlsx) as an engagement.

The sheet is the base: its project block, staffing table and additional-costs table carry everything
needed to create the engagement, its allocations and its costs. Columns are found by header text,
not position, so older copies of the sheet import too.

parse_workbook(path_or_file) -> ImportedEngagement   (pure; no DB, no engine call)
The API layer creates the rows and returns compute() of the result plus the `warnings` here.
"""
from __future__ import annotations
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Optional
import re

from openpyxl import load_workbook

D = Decimal
TYPE_MAP = {"fixed cost": "fixed_cost", "assessment": "assessment", "msp": "msp",
            "time & materials": "time_and_materials", "time and materials": "time_and_materials", "t&m": "time_and_materials",
            "staff augmentation": "staff_augmentation", "staff aug": "staff_augmentation"}
EMP_MAP = {"w-2": "w2_employee", "w2": "w2_employee", "employee": "w2_employee", "1099": "contractor_1099", "contractor": "contractor_1099"}


@dataclass
class ImportedAllocation:
    role: str
    seniority: Optional[str]
    location: str
    employment_type: str
    assignment: Optional[str]
    utilization: Decimal
    bill_rate: Optional[Decimal]
    cost_per_hour: Optional[Decimal]
    start: Optional[date]
    end: Optional[date]
    hours_override: Optional[Decimal]
    sheet_hours: Optional[Decimal]          # what the sheet showed — the API reports engine hours beside it


@dataclass
class ImportedCost:
    description: str
    category: Optional[str]
    amount: Decimal
    frequency: str                          # one_time | monthly
    billed_to_client: bool


@dataclass
class ImportedEngagement:
    client: str
    name: str
    type: str
    currency: str
    start: date
    end: date
    total_price: Optional[Decimal]
    monthly_fee: Optional[Decimal]
    additional_costs_flag: bool
    same_dates_for_all: bool
    allocations: list = field(default_factory=list)
    costs: list = field(default_factory=list)
    sheet_settings: dict = field(default_factory=dict)   # burden / floors / hours per day as the sheet had them, for comparison
    warnings: list = field(default_factory=list)
    errors: list = field(default_factory=list)


# ----------------------------------------------------------------------------- helpers
def _txt(v) -> str:
    return "" if v is None else str(v).strip()

def _lower(v) -> str:
    return _txt(v).lower()

def _dec(v) -> Optional[Decimal]:
    if v is None or _txt(v) == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return D(str(v))
    t = _txt(v).replace("$", "").replace(",", "").replace("%", "")
    try:
        return D(t)
    except Exception:
        return None

def _date(v) -> Optional[date]:
    if v is None or _txt(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _txt(v)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%b %d, %Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None

def _pct(v) -> Optional[Decimal]:
    x = _dec(v)
    if x is None:
        return None
    return x / D(100) if x > 1 else x

def _find_label(ws, label_prefix: str, max_row: int = 40):
    """Return the value in column B of the first row whose column A starts with label_prefix (case-insensitive)."""
    lp = label_prefix.lower()
    for r in range(1, max_row + 1):
        if _lower(ws.cell(row=r, column=1).value).startswith(lp):
            return ws.cell(row=r, column=2).value
    return None

def _header_row(ws, must_contain: str, max_row: int = 60) -> Optional[int]:
    for r in range(1, max_row + 1):
        for c in range(1, 30):
            if _lower(ws.cell(row=r, column=c).value).startswith(must_contain):
                return r
    return None

def _columns(ws, row: int) -> dict:
    """Map a header prefix (lower-case) to its column index."""
    out = {}
    for c in range(1, 40):
        h = _lower(ws.cell(row=row, column=c).value)
        if h:
            out[h] = c
    return out

def _col(cols: dict, *prefixes) -> Optional[int]:
    for p in prefixes:
        for h, c in cols.items():
            if h.startswith(p):
                return c
    return None


# ----------------------------------------------------------------------------- parser
def parse_workbook(src) -> ImportedEngagement:
    wb = load_workbook(src, data_only=True)
    if "Cost Margin" not in wb.sheetnames:
        raise ValueError("This is not a DealGate cost-margin workbook: no 'Cost Margin' sheet.")
    ws = wb["Cost Margin"]
    warnings, errors = [], []

    # --- project block (labels in column A, values in column B) ---
    typ_raw = _lower(_find_label(ws, "engagement type"))
    typ = TYPE_MAP.get(typ_raw)
    if not typ:
        errors.append(f"Engagement type '{typ_raw or '(blank)'}' not recognised — pick one on the sheet first.")
    name = _txt(_find_label(ws, "project / engagement name") or _find_label(ws, "project name"))
    client = _txt(_find_label(ws, "client"))
    currency = _txt(_find_label(ws, "currency")) or "USD"
    start = _date(_find_label(ws, "project start")); end = _date(_find_label(ws, "project end"))
    if not start or not end:
        errors.append("Project start and end dates are required.")
    elif end < start:
        errors.append("Project end is before start.")
    price = _dec(_find_label(ws, "total project price")); fee = _dec(_find_label(ws, "monthly fee"))
    addl_flag = _lower(_find_label(ws, "additional costs on this project")).startswith("y")
    same_raw = _lower(_find_label(ws, "same start/end"))
    same = same_raw.startswith("y") if same_raw else False
    if not client:
        errors.append("Client is blank on the sheet.")
    if not name:
        name = f"{client} engagement" if client else "Imported engagement"
        warnings.append("Project name was blank; used a placeholder — rename it.")
    if typ in ("fixed_cost", "assessment") and price is None:
        errors.append("Total project price is required for Fixed cost / Assessment.")
    if typ == "msp" and fee is None:
        errors.append("Monthly fee is required for MSP.")

    # --- staffing table ---
    hr = _header_row(ws, "role")
    allocations = []
    if hr is None:
        errors.append("Staffing table not found (no 'Role' header).")
    else:
        cols = _columns(ws, hr)
        c_role = _col(cols, "role"); c_sen = _col(cols, "seniority"); c_loc = _col(cols, "location")
        c_emp = _col(cols, "employment"); c_asg = _col(cols, "assignment"); c_util = _col(cols, "utilization")
        c_hours = _col(cols, "hours  (auto", "hours (auto", "hours (calculated", "hours *", "hours")
        c_bill = _col(cols, "bill rate"); c_cost = _col(cols, "our cost", "cost / hour", "cost/hour")
        c_start = _col(cols, "start"); c_end = _col(cols, "end"); c_ov = _col(cols, "hours override")
        if c_emp is None:
            warnings.append("This copy of the sheet has no Employment type column; every person was imported as W-2 / Employee — fix any 1099s on the Allocations tab.")
        for r in range(hr + 1, hr + 60):
            role = _txt(ws.cell(row=r, column=c_role).value) if c_role else ""
            if not role:
                continue
            if role.lower().startswith("staffing total"):
                break
            loc = _txt(ws.cell(row=r, column=c_loc).value) if c_loc else ""
            emp_raw = _lower(ws.cell(row=r, column=c_emp).value) if c_emp else ""
            emp = next((v for k, v in EMP_MAP.items() if k in emp_raw), "w2_employee")
            util = _pct(ws.cell(row=r, column=c_util).value) if c_util else None
            a = ImportedAllocation(
                role=role, seniority=_txt(ws.cell(row=r, column=c_sen).value) or None if c_sen else None,
                location=loc or "US", employment_type=emp,
                assignment=(_txt(ws.cell(row=r, column=c_asg).value) or None) if c_asg else None,
                utilization=util if util is not None else D("1"),
                bill_rate=_dec(ws.cell(row=r, column=c_bill).value) if c_bill else None,
                cost_per_hour=_dec(ws.cell(row=r, column=c_cost).value) if c_cost else None,
                start=_date(ws.cell(row=r, column=c_start).value) if c_start else None,
                end=_date(ws.cell(row=r, column=c_end).value) if c_end else None,
                hours_override=_dec(ws.cell(row=r, column=c_ov).value) if c_ov else None,
                sheet_hours=_dec(ws.cell(row=r, column=c_hours).value) if c_hours else None)
            if not loc:
                warnings.append(f"{role}: location blank on the sheet; assumed US.")
            if a.start == start and a.end == end:
                a.start = a.end = None          # same as the window: let the engine default them
            if typ in ("time_and_materials", "staff_augmentation") and a.bill_rate is None:
                errors.append(f"{role}: bill rate is required for {typ.replace('_', ' ')}.")
            allocations.append(a)
        if not allocations:
            errors.append("No staffing rows found.")

    # --- additional costs ---
    costs = []
    cr = _header_row(ws, "description")
    if cr is not None:
        cols = _columns(ws, cr)
        c_desc = _col(cols, "description"); c_cat = _col(cols, "category"); c_amt = _col(cols, "amount")
        c_freq = _col(cols, "frequency"); c_billed = _col(cols, "billed to client")
        for r in range(cr + 1, cr + 40):
            desc = _txt(ws.cell(row=r, column=c_desc).value) if c_desc else ""
            if not desc:
                continue
            if desc.lower().startswith("additional costs total"):
                break
            amt = _dec(ws.cell(row=r, column=c_amt).value) if c_amt else None
            if amt is None:
                warnings.append(f"Cost '{desc}' has no amount; skipped.")
                continue
            freq = "monthly" if (c_freq and _lower(ws.cell(row=r, column=c_freq).value).startswith("month")) else "one_time"
            billed_raw = _lower(ws.cell(row=r, column=c_billed).value) if c_billed else ""
            if addl_flag and not billed_raw:
                errors.append(f"Cost '{desc}': 'Billed to client' Y/N is blank on the sheet.")
            costs.append(ImportedCost(desc, _txt(ws.cell(row=r, column=c_cat).value) or None if c_cat else None, amt, freq, billed_raw.startswith("y")))
    if addl_flag and not costs:
        warnings.append("Additional costs is Y but the table is empty.")
    if not addl_flag and costs:
        warnings.append("Additional costs is N; the cost rows on the sheet were imported but will not count until you set it to Y.")

    # --- settings, for comparison only (the app's Settings page is the source of truth) ---
    settings = {}
    if "Settings" in wb.sheetnames:
        st = wb["Settings"]
        hpd = _dec(_find_label(st, "hours per working day", 10))
        if hpd:
            settings["hours_per_day"] = hpd
        pf = _dec(_find_label(st, "project gm floor", 10))
        if pf is not None:
            settings["project_floor_pct"] = pf if pf <= 1 else pf / 100
        lr = _header_row(st, "location", 12)
        if lr:
            cols = _columns(st, lr); c_l = _col(cols, "location"); c_w2 = _col(cols, "burden % — w-2", "burden %"); c_99 = _col(cols, "burden % — 1099"); c_fl = _col(cols, "resource gm floor", "target gm", "gm floor")
            for r in range(lr + 1, lr + 6):
                L = _txt(st.cell(row=r, column=c_l).value) if c_l else ""
                if not L:
                    continue
                settings.setdefault("locations", {})[L] = {
                    "burden_w2": _pct(st.cell(row=r, column=c_w2).value) if c_w2 else None,
                    "burden_1099": _pct(st.cell(row=r, column=c_99).value) if c_99 else (_pct(st.cell(row=r, column=c_w2).value) if c_w2 else None),
                    "resource_floor": _pct(st.cell(row=r, column=c_fl).value) if c_fl else None}
        if settings:
            warnings.append("Settings on the sheet (burden, floors, hours/day) were read for comparison only; the app's Settings page is used for the forecast.")

    return ImportedEngagement(client, name, typ or "", currency, start, end, price, fee, addl_flag, same, allocations, costs, settings, warnings, errors)
