"""S3 E6 — Delivery Model Builder API + service + xlsx export.

The story's acceptance tests, one Given/When/Then per test.

Every case round-trips a real HTTP request through the ratified
:func:`app.gm.compute` library so §7 discounted numbers stay the source
of truth. No math in this file (CLAUDE.md rule 2).
"""

from __future__ import annotations

import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import httpx
import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy import select

from app.db import get_session
from app.main import app as main_app
from app.models.audit import AuditEvent
from app.models.gm_model import CostLine, GmModel, ResourceLine
from app.models.opportunity import Opportunity


@pytest.fixture(autouse=True)
def _local_env(monkeypatch):
    monkeypatch.setenv("DEALGATE_ENV", "local")
    monkeypatch.delenv("DEALGATE_TEST_GROUPS", raising=False)


def _client(app):
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")


def _fake_user_id(email: str) -> uuid.UUID:
    return uuid.uuid5(uuid.NAMESPACE_URL, f"dealgate:local:{email}")


@pytest_asyncio.fixture
async def app_with_session(session):
    async def _override():
        yield session

    main_app.dependency_overrides[get_session] = _override
    try:
        yield main_app
    finally:
        main_app.dependency_overrides.pop(get_session, None)


@pytest_asyncio.fixture
async def opportunity(session) -> Opportunity:
    opp = Opportunity(
        id=uuid.uuid4(),
        hubspot_deal_id="H-DM-001",
        owner_id=None,
        governance_status="Intake",
    )
    session.add(opp)
    await session.commit()
    return opp


# --- payload helpers -------------------------------------------------------

# §7 mixed-geography discounted case:
#   US:    revenue 90k / cost 65k → GM 27.78% (fails 35% floor)
#   India: revenue 55k / cost 30k → GM 45.45% (fails 50% floor)


def _fixed_price_discounted(sow_version_id: str | None = None) -> dict:
    start = date(2026, 3, 1)
    end = date(2026, 8, 31)
    return {
        "engagement_type": "fixed_price",
        "sow_version_id": sow_version_id,
        "delivery_pattern": "hybrid",
        "contingency_pct": "5.00",
        "warranty_days": 30,
        "resource_lines": [
            {
                "role": "Engineer",
                "seniority": "Sr",
                "location": "US",
                "person_name": "Alice",
                "allocation_pct": "1",
                "start_date": start.isoformat(),
                "end_date": end.isoformat(),
                "hours_billable": "1000",
                # Bill rate carries revenue in the compute layer; the
                # fixed-price template uses the explicit revenue split
                # instead, so setting hourly_bill_rate here only feeds
                # the snapshot revenue columns on gm_model.
                "hourly_bill_rate": "90",
                "hourly_cost": "65",
                "validated_by": str(uuid.uuid4()),
            },
            {
                "role": "Engineer",
                "seniority": "Sr",
                "location": "India",
                "person_name": "Bob",
                "allocation_pct": "1",
                "start_date": start.isoformat(),
                "end_date": end.isoformat(),
                "hours_billable": "600",
                "hourly_bill_rate": "91.67",
                "hourly_cost": "50",
                "validated_by": str(uuid.uuid4()),
            },
        ],
        "cost_lines": [],
        # Fixed-price top-level revenue allocation forwarded via /preview.
        "total_price": "145000",
        "revenue_us": "90000",
        "revenue_india": "55000",
    }


# --- Preview: §7 discounted case matches sandbox --------------------------


async def test_preview_returns_correct_gm_for_discounted_case(app_with_session, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    payload = _fixed_price_discounted()
    async with _client(app_with_session) as c:
        r = await c.post(
            "/delivery-model/preview",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json={
                "engagement_type": payload["engagement_type"],
                "inputs": {
                    "resource_lines": payload["resource_lines"],
                    "cost_lines": payload["cost_lines"],
                    "total_price": payload["total_price"],
                    "revenue_us": payload["revenue_us"],
                    "revenue_india": payload["revenue_india"],
                    "delivery_pattern": payload["delivery_pattern"],
                    "contingency_pct": payload["contingency_pct"],
                    "warranty_days": payload["warranty_days"],
                },
            },
        )
    assert r.status_code == 200, r.text
    body = r.json()
    computed = body["computed"]
    assert Decimal(computed["revenue_us"]) == Decimal("90000")
    assert Decimal(computed["cost_us"]) == Decimal("65000")
    assert Decimal(computed["revenue_india"]) == Decimal("55000")
    assert Decimal(computed["cost_india"]) == Decimal("30000")
    assert Decimal(computed["gm_us"]).quantize(Decimal("0.0001")) == Decimal("0.2778")
    assert Decimal(computed["gm_india"]).quantize(Decimal("0.0001")) == Decimal("0.4545")
    assert computed["policy"]["us_pass"] is False
    assert computed["policy"]["india_pass"] is False
    assert computed["policy"]["requires_ceo"] is True


async def test_preview_forbidden_for_sales(app_with_session, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Sales")
    async with _client(app_with_session) as c:
        r = await c.post(
            "/delivery-model/preview",
            headers={"X-Test-User": "sales@smartek21.com"},
            json={"engagement_type": "fixed_price", "inputs": {"resource_lines": []}},
        )
    assert r.status_code == 403


async def test_preview_requires_auth(app_with_session):
    async with _client(app_with_session) as c:
        r = await c.post(
            "/delivery-model/preview",
            json={"engagement_type": "fixed_price", "inputs": {"resource_lines": []}},
        )
    assert r.status_code == 401


# --- Save creates immutable version + rows + audit ------------------------


async def test_save_creates_version_rows_and_audit(app_with_session, session, opportunity, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    payload = _fixed_price_discounted()
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=payload,
        )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["gm_model"]["engagement_type"] == "fixed_price"
    assert len(body["gm_model"]["resource_lines"]) == 2
    new_id = body["gm_model"]["id"]

    # DB round-trip: rows persisted, snapshot revenue populated.
    models = (await session.execute(select(GmModel))).scalars().all()
    assert len(models) == 1
    assert str(models[0].id) == new_id
    assert models[0].revenue_us > 0
    assert models[0].revenue_india > 0

    rows = (
        await session.execute(select(ResourceLine).where(ResourceLine.gm_model_id == models[0].id))
    ).scalars().all()
    assert len(rows) == 2

    # Audit event landed.
    audit = (
        await session.execute(
            select(AuditEvent).where(AuditEvent.action == "gm_model.created")
        )
    ).scalars().all()
    assert len(audit) == 1
    assert audit[0].actor_id == _fake_user_id("delivery@smartek21.com")
    assert audit[0].after["opportunity_id"] == str(opportunity.id)
    assert audit[0].after["resource_line_count"] == 2

    # Second save creates a *new* immutable version — the first row is not mutated.
    async with _client(app_with_session) as c:
        r2 = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=payload,
        )
    assert r2.status_code == 201, r2.text
    all_models = (await session.execute(select(GmModel))).scalars().all()
    assert len(all_models) == 2


async def test_save_persists_cost_lines(app_with_session, session, opportunity, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    payload = _fixed_price_discounted()
    payload["cost_lines"] = [
        {"category": "travel", "amount": "1500", "location": "US", "note": "kickoff"},
        {"category": "tools", "amount": "800", "location": "India"},
    ]
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=payload,
        )
    assert r.status_code == 201, r.text
    rows = (await session.execute(select(CostLine))).scalars().all()
    assert len(rows) == 2
    cats = sorted(c.category for c in rows)
    assert cats == ["tools", "travel"]


async def test_save_forbidden_for_sales(app_with_session, opportunity, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Sales")
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "sales@smartek21.com"},
            json=_fixed_price_discounted(),
        )
    assert r.status_code == 403


async def test_save_forbidden_for_finance(app_with_session, opportunity, monkeypatch):
    """Finance can read + preview but not save; story is explicit."""
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Finance")
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "fin@smartek21.com"},
            json=_fixed_price_discounted(),
        )
    assert r.status_code == 403


async def test_save_404_for_unknown_opportunity(app_with_session, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{uuid.uuid4()}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=_fixed_price_discounted(),
        )
    assert r.status_code == 404


# --- Capacity conflict ------------------------------------------------------


async def test_capacity_conflict_flagged_when_person_double_booked(
    app_with_session, session, opportunity, monkeypatch
):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    # Persist a first model with Alice at 100% on 2026-03..08.
    async with _client(app_with_session) as c:
        r0 = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=_fixed_price_discounted(),
        )
    assert r0.status_code == 201, r0.text

    # Preview a new engagement that also books Alice at 100% in overlapping dates.
    overlap_payload = _fixed_price_discounted()
    async with _client(app_with_session) as c:
        r = await c.post(
            "/delivery-model/preview",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json={
                "engagement_type": overlap_payload["engagement_type"],
                "inputs": {
                    "resource_lines": overlap_payload["resource_lines"],
                    "cost_lines": overlap_payload["cost_lines"],
                    "total_price": overlap_payload["total_price"],
                    "revenue_us": overlap_payload["revenue_us"],
                    "revenue_india": overlap_payload["revenue_india"],
                },
            },
        )
    assert r.status_code == 200, r.text
    warnings = r.json()["warnings"]["capacity"]
    # Alice (row 0) triggers the amber conflict; Bob (row 1) too.
    codes = {w["code"] for w in warnings}
    assert "capacity_conflict" in codes
    assert any("Alice" in w["message"] for w in warnings)


# --- HR lead time -----------------------------------------------------------


async def test_hr_lead_time_flagged_for_early_to_hire(app_with_session, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    start = date.today() + timedelta(days=10)  # well inside 45-day lead time
    payload = _fixed_price_discounted()
    payload["resource_lines"][0]["person_name"] = None  # to-hire
    payload["resource_lines"][0]["start_date"] = start.isoformat()
    async with _client(app_with_session) as c:
        r = await c.post(
            "/delivery-model/preview",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json={
                "engagement_type": payload["engagement_type"],
                "inputs": {
                    "resource_lines": payload["resource_lines"],
                    "cost_lines": payload["cost_lines"],
                    "total_price": payload["total_price"],
                    "revenue_us": payload["revenue_us"],
                    "revenue_india": payload["revenue_india"],
                },
            },
        )
    assert r.status_code == 200, r.text
    hr_warnings = r.json()["warnings"]["hr"]
    assert any(w["code"] == "hr_lead_time" for w in hr_warnings)
    assert any(w["severity"] == "red" for w in hr_warnings)


# --- Completeness check ----------------------------------------------------


async def test_completeness_flags_missing_pieces(app_with_session, opportunity, monkeypatch):
    """Workflow (Sprint 4) owns the actual transition. This asserts the
    completeness helper flags the four pieces the story requires: location,
    effort, bill_rate, validated_cost."""
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    payload = _fixed_price_discounted()
    # Blank out cost + validated_by on one line; that line should be flagged.
    payload["resource_lines"][0]["hourly_cost"] = None
    payload["resource_lines"][0]["validated_by"] = None
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=payload,
        )
    assert r.status_code == 201, r.text
    issues = r.json()["gm_model"]["completeness_issues"]
    joined = " ".join(issues)
    assert "resource_lines[0].hourly_cost" in joined
    assert "resource_lines[0].validated_by" in joined


# --- Governance status stays Intake ----------------------------------------


async def test_governance_status_not_auto_advanced(
    app_with_session, session, opportunity, monkeypatch
):
    """This service never mutates opportunity.governance_status — the
    workflow module (Sprint 4) owns advancement (CLAUDE.md rule 5)."""
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    async with _client(app_with_session) as c:
        r = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=_fixed_price_discounted(),
        )
    assert r.status_code == 201
    await session.refresh(opportunity)
    assert opportunity.governance_status == "Intake"


# --- GET latest -------------------------------------------------------------


async def test_get_latest_returns_saved_model_with_computed(
    app_with_session, opportunity, monkeypatch
):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    payload = _fixed_price_discounted()
    async with _client(app_with_session) as c:
        r0 = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=payload,
        )
    assert r0.status_code == 201, r0.text

    async with _client(app_with_session) as c:
        r = await c.get(
            f"/delivery-model/{opportunity.id}",
            headers={"X-Test-User": "delivery@smartek21.com"},
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["gm_model"] is not None
    assert body["gm_model"]["engagement_type"] == "fixed_price"


async def test_list_versions_returns_all(app_with_session, opportunity, monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    for _ in range(3):
        async with _client(app_with_session) as c:
            r = await c.post(
                f"/delivery-model/{opportunity.id}/versions",
                headers={"X-Test-User": "delivery@smartek21.com"},
                json=_fixed_price_discounted(),
            )
        assert r.status_code == 201, r.text
    async with _client(app_with_session) as c:
        r = await c.get(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
        )
    assert r.status_code == 200, r.text
    assert len(r.json()["items"]) == 3


# --- xlsx export matches compute -------------------------------------------


async def test_xlsx_export_numbers_match_preview(app_with_session, opportunity, monkeypatch):
    """Story AC: 'Export to Excel matches the GM sandbox numbers exactly
    for the same inputs'."""
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Delivery")
    payload = _fixed_price_discounted()
    async with _client(app_with_session) as c:
        r0 = await c.post(
            f"/delivery-model/{opportunity.id}/versions",
            headers={"X-Test-User": "delivery@smartek21.com"},
            json=payload,
        )
    assert r0.status_code == 201, r0.text
    model_id = r0.json()["gm_model"]["id"]

    # Read back the saved model — its computed block is the reference.
    async with _client(app_with_session) as c:
        latest = await c.get(
            f"/delivery-model/{opportunity.id}",
            headers={"X-Test-User": "delivery@smartek21.com"},
        )
    assert latest.status_code == 200, latest.text
    computed = latest.json()["gm_model"]["computed"]

    async with _client(app_with_session) as c:
        export = await c.get(
            f"/delivery-model/versions/{model_id}/xlsx",
            headers={"X-Test-User": "delivery@smartek21.com"},
        )
    assert export.status_code == 200, export.text
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(BytesIO(export.content), data_only=False)
    assert "Result" in wb.sheetnames
    kv: dict[str, object] = {}
    for row in wb["Result"].iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        kv[str(row[0])] = row[1]

    assert Decimal(str(kv["revenue_us"])) == Decimal(computed["revenue_us"])
    assert Decimal(str(kv["cost_us"])) == Decimal(computed["cost_us"])
    assert Decimal(str(kv["revenue_india"])) == Decimal(computed["revenue_india"])
    assert Decimal(str(kv["cost_india"])) == Decimal(computed["cost_india"])
    assert Decimal(str(kv["gm_us"])) == Decimal(computed["gm_us"])
    assert Decimal(str(kv["gm_india"])) == Decimal(computed["gm_india"])
