"""S22 P5 — business units, the forecast-model master, actuals (directive §P5 items 2–4, 9).

Drives the real router over HTTP against the SQLite test schema, like test_s22_portfolio_forecast.
The model fixture is the synthetic workbook test_forecast_model_import builds (same layouts as the
real model, invented people and clients); the real model never enters the repo.
"""

from __future__ import annotations

import io
from decimal import Decimal as D
from pathlib import Path

import pytest
from sqlalchemy import func, select

from app.models.client import Client
from app.models.engagement import Allocation, Engagement
from tests.forecast.test_forecast_model_import import build_fixture
from tests.test_s22_portfolio_forecast import (
    actor_for,
    client_for,
    close,
    post,
    seed_calendar,
    seed_client,
    seed_client_a,
)

MODEL = Path(__file__).parent / "forecast" / "fixtures" / "forecast_model_synthetic.xlsx"
# The same workbook without the "Zeta Unlisted" tab (no dashboard row, one blank location): nothing blocks a load.
# Playwright uses both files: this one for the owner's load path, the full one for the blockers panel (P5c).
COMPLETE = Path(__file__).parent / "forecast" / "fixtures" / "forecast_model_synthetic_complete.xlsx"


def model_bytes(complete: bool = False) -> bytes:
    if not MODEL.exists():
        build_fixture(MODEL)
    if complete:
        import openpyxl

        wb = openpyxl.load_workbook(MODEL)
        if "Zeta Unlisted" in wb.sheetnames:
            wb.remove(wb["Zeta Unlisted"])
        wb.save(COMPLETE)
        return COMPLETE.read_bytes()
    return MODEL.read_bytes()


FIXED = Path(__file__).parent / "forecast" / "fixtures" / "forecast_model_synthetic_fixed.xlsx"


def fixed_bytes() -> bytes:
    """The full fixture with its two blockers fixed in the workbook, the way the product owner will: Zeta Unlisted
    gets a dashboard row under ST21 and its blank Location is filled. Nothing else changes — the two TBC seats and
    the extension cases stay. Built on demand, never committed."""
    import openpyxl

    model_bytes()
    wb = openpyxl.load_workbook(MODEL)
    dash = wb["Dashboard (RunRate)"]
    subtotal_row = next(r for r in range(6, 40) if str(dash.cell(row=r, column=2).value or "").strip().lower() == "subtotal")
    dash.insert_rows(subtotal_row)
    for col, value in enumerate(("Zeta Unlisted", "Sam", "Dee", "T&M"), start=2):
        dash.cell(row=subtotal_row, column=col, value=value)
    zeta = wb["Zeta Unlisted"]
    headers = {str(zeta.cell(row=4, column=c).value or "").strip().lower(): c for c in range(1, 20)}
    loc_col = headers["location"]
    for r in range(5, 12):
        if zeta.cell(row=r, column=1).value and str(zeta.cell(row=r, column=1).value).strip().lower() != "total" and not zeta.cell(row=r, column=loc_col).value:
            zeta.cell(row=r, column=loc_col, value="US")
    wb.save(FIXED)
    return FIXED.read_bytes()


async def upload_model(client, *, dry_run: bool, as_of: str | None = "2026-02", expect: int = 200, complete: bool = False, fixed: bool = False) -> dict:
    data = fixed_bytes() if fixed else model_bytes(complete)
    response = await client.post(
        "/engagements/import-forecast-model",
        files={"file": ("model.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"dry_run": "true" if dry_run else "false", **({"as_of": as_of} if as_of else {})},
    )
    assert response.status_code == expect, response.text
    return response.json()


async def test_model_import_dry_run_reports_three_business_units_and_creates_nothing(session):
    await seed_calendar(session)
    actor = await actor_for(session, ("SystemAdmin",))
    async with client_for(session, actor) as client:
        report = await upload_model(client, dry_run=True)
        assert report["dry_run"] is True and report["created"] is None
        assert [b["business_unit"] for b in report["by_business_unit"]] == ["ST21", "IT Avalon", "Retro Rabbit", "Unassigned"]
        st21 = next(b for b in report["by_business_unit"] if b["business_unit"] == "ST21")
        assert st21["clients"] == 2 and st21["engagements"] == 2 and st21["people"] == 5
        avalon = next(b for b in report["by_business_unit"] if b["business_unit"] == "IT Avalon")
        assert avalon["engagements"] == 3 and avalon["closed"] == 1            # Gamma VMS + PO, Delta (closed)
        assert {t["tab"] for t in report["tabs"]} == {"Alpha", "Beta", "Gamma Health-VMS", "Delta Closed", "Zeta Unlisted"}
        assert {(b["kind"], b["tab"]) for b in report["blockers"]} == {("business_unit", "Zeta Unlisted"), ("location", "Zeta Unlisted")}
        assert "## By business unit" in report["report_markdown"] and "## Calendar" in report["report_markdown"]
        assert "India - Pune" in report["report_markdown"] and "South Africa" in report["report_markdown"]
        assert (await client.get("/engagements")).json()["total"] == 0
        assert await session.scalar(select(func.count()).select_from(Engagement)) == 0


async def test_model_import_load_creates_the_ledger_and_the_portfolio_groups_reconcile(session):
    await seed_calendar(session)
    actor = await actor_for(session, ("SystemAdmin",))
    async with client_for(session, actor) as client:
        loaded = await upload_model(client, dry_run=False, complete=True)
        assert loaded["blockers"] == []
        created = loaded["created"]
        assert created["clients"] == 4 and created["engagements"] == 5 and created["allocations"] >= 8 and created["actual_rows"] > 0
        # locations the Settings page lacked now exist with their WorkingDays holidays
        settings = (await client.get("/settings/forecast")).json()
        names = {loc["name"] for loc in settings["locations"]}
        assert {"India - Pune", "India - Chennai", "South Africa"} <= names
        assert len(settings["holidays"]["India - Pune"]) == 4 and len(settings["holidays"]["South Africa"]) == 1
        # closed engagements are kept, flagged, and off the portfolio
        listed = (await client.get("/engagements?page_size=50")).json()["items"]
        delta = next(e for e in listed if e["name"].startswith("Delta Closed"))
        assert delta["closed"] is True and delta["business_unit"] == "IT Avalon"
        gamma = [e for e in listed if e["name"].startswith("Gamma Health-VMS")]
        assert len(gamma) == 2 and all(e["business_unit"] == "IT Avalon" and e["sales_owner"] == "Kim" and e["delivery_owner"] == "Lee" for e in gamma)
        # imported people have no cost: INCOMPLETE until the rate card covers them; actuals arrived with the sheet
        alpha = next(e for e in listed if e["name"].startswith("Alpha"))
        f = (await client.get(f"/engagements/{alpha['id']}/forecast")).json()
        assert f["status"] == "INCOMPLETE" and any("cost per hour unknown" in m for m in f["missing"])
        assert f["actuals"]["months_with_actuals"] == 3
        sep = next(m for m in f["months"] if m["month"] == "2025-09")
        assert sep["actual_revenue"] is not None and close(sep["billing_loss"], D(sep["revenue"]) - D(sep["actual_revenue"]))
        assert close(f["actuals"]["billing_loss"], D(f["actuals"]["forecast_in_actual_months"]) - D(f["actuals"]["actual_revenue"]))
        rows = (await client.get(f"/engagements/{alpha['id']}/allocations")).json()
        assert {r["person"] for r in rows} == {"Person A1", "Person A2", "Person A3"} and all(r["rate_basis"] == "monthly" for r in rows)
        assert all(r["cost_per_hour"] is None for r in rows)

        # item 9: BU subtotals == Σ their rows; grand total == Σ subtotals; billing-loss count == rows with actuals
        month = "2026-01"
        view = (await client.get(f"/portfolio/forecast?month={month}")).json()
        assert delta["id"] not in {r["engagement_id"] for r in view["rows"]}
        assert [g["business_unit"] for g in view["groups"]] == ["ST21", "IT Avalon"]
        for g in view["groups"]:
            mine = [r for r in view["rows"] if r["business_unit"] == g["business_unit"]]
            assert g["engagements"] == len(mine) and g["people"] == sum(r["people"] for r in mine)
            assert close(g["revenue"], sum(D(r["revenue_month"]) for r in mine))
            assert close(g["cost"], sum(D(r["cost_month"]) for r in mine))
        assert close(view["totals"]["revenue"], sum(D(g["revenue"]) for g in view["groups"]))
        assert close(view["totals"]["cost"], sum(D(g["cost"]) for g in view["groups"]))
        assert view["totals"]["people"] == sum(g["people"] for g in view["groups"])
        with_actuals = 0
        for r in view["rows"]:
            fr = (await client.get(f"/engagements/{r['engagement_id']}/forecast")).json()
            with_actuals += 1 if fr["actuals"]["months_with_actuals"] else 0
            assert r["sales_owner"] and r["delivery_owner"]
        assert view["cards"]["billing_loss"]["engagements"] == with_actuals > 0
        assert close(view["cards"]["billing_loss"]["amount"], D(view["cards"]["billing_loss"]["forecast"]) - D(view["cards"]["billing_loss"]["billed"]))
        # one BU tab: its rows and one subtotal only
        tab = (await client.get(f"/portfolio/forecast?month={month}&business_unit=IT%20Avalon")).json()
        assert {r["business_unit"] for r in tab["rows"]} == {"IT Avalon"} and [g["business_unit"] for g in tab["groups"]] == ["IT Avalon"]
        assert close(tab["totals"]["revenue"], tab["groups"][0]["revenue"])
        csv_text = (await client.get(f"/portfolio/forecast.csv?month={month}")).text
        assert "ST21 subtotal" in csv_text and "IT Avalon subtotal" in csv_text and f"All business units · {month}" in csv_text

        # re-importing the same file is idempotent
        again = await upload_model(client, dry_run=False, complete=True)
        assert again["created"]["engagements"] == 0 and again["created"]["allocations"] == 0 and again["created"]["clients"] == 0
        assert (await client.get("/engagements?page_size=50")).json()["total"] == 5
        assert len((await client.get(f"/engagements/{alpha['id']}/allocations")).json()) == 3


async def test_put_actuals_recomputes_the_forecast_and_the_billing_loss_card(session):
    await seed_calendar(session)
    actor = await actor_for(session, ("Delivery",))
    client_row = await seed_client(session, "Client A")
    async with client_for(session, actor) as client:
        eng, allocations = await seed_client_a(client, client_row)
        before = (await client.get(f"/engagements/{eng['id']}/forecast")).json()
        assert before["actuals"]["months_with_actuals"] == 0 and all(m["actual_revenue"] is None for m in before["months"])
        first_month = before["months"][0]["month"]
        response = await client.put(
            f"/allocations/{allocations[0]['id']}/actuals",
            json=[{"month": first_month, "hours": 150, "revenue": "20000.00"}],
        )
        assert response.status_code == 200, response.text
        f = response.json()
        assert f["actuals"]["months_with_actuals"] == 1
        line = next(m for m in f["months"] if m["month"] == first_month)
        assert close(line["actual_revenue"], "20000.00") and close(line["actual_hours"], 150)
        assert close(line["billing_loss"], D(line["revenue"]) - D("20000.00"))
        assert close(f["actuals"]["billing_loss"], D(line["revenue"]) - D("20000.00"))
        # the plan numbers are untouched by actuals
        assert f["totals"] == before["totals"]
        view = (await client.get(f"/portfolio/forecast?month={first_month}")).json()
        assert view["cards"]["billing_loss"]["engagements"] == 1 and close(view["cards"]["billing_loss"]["amount"], f["actuals"]["billing_loss"])
        # months not in the body are removed
        cleared = (await client.put(f"/allocations/{allocations[0]['id']}/actuals", json=[])).json()
        assert cleared["actuals"]["months_with_actuals"] == 0
        bad = await client.put(f"/allocations/{allocations[0]['id']}/actuals", json=[{"month": "2026-13"}])
        assert bad.status_code == 422


async def test_business_unit_owners_closed_rate_basis_and_msp_on_monthly_rates(session):
    await seed_calendar(session)
    actor = await actor_for(session, ("Delivery",))
    client_row = await seed_client(session, "Client C")
    async with client_for(session, actor) as client:
        eng = await post(client, "/engagements", {
            "client_id": str(client_row.id), "name": "O365 support — CR #01", "type": "msp", "start": "2026-07-01", "end": "2027-06-30",
            "additional_costs_flag": False, "same_dates_for_all": True, "business_unit": "IT Avalon", "sales_owner": "Janice Krpan",
            "delivery_owner": "Srikanth Parimi",
        })
        assert eng["business_unit"] == "IT Avalon" and eng["sales_owner"] == "Janice Krpan" and eng["closed"] is False
        await post(client, "/engagements", {
            "client_id": str(client_row.id), "name": "bad", "type": "msp", "start": "2026-07-01", "end": "2026-12-31",
            "additional_costs_flag": False, "same_dates_for_all": True, "business_unit": "Nope", "monthly_fee": 1,
        }, expect=422)
        res = await post(client, "/resources", {"name": "Person 1", "location": "India", "employment_type": "w2_employee", "seniority": "Mid"})
        row = await post(client, f"/engagements/{eng['id']}/allocations", {
            "resource_id": res["id"], "role": "Developer", "utilization": 1, "bill_rate": 5151, "rate_basis": "monthly", "person": "Person 1",
        })
        assert row["rate_basis"] == "monthly" and row["person"] == "Person 1"
        assert D(row["totals"]["revenue"]) > 0 and row["totals"]["row_gm_pct"] is not None      # MSP without a fee bills the people's rates
        f = (await client.get(f"/engagements/{eng['id']}/forecast")).json()
        assert not any("monthly_fee" in m for m in f["missing"]) and f["status"] in ("PASS", "BELOW FLOOR")
        jul = next(m for m in f["months"] if m["month"] == "2026-07")
        assert close(jul["revenue"], "5151.00")                                                   # a monthly rate is calendar-independent
        patched = await client.patch(f"/allocations/{row['id']}", json={"rate_basis": "weekly"})
        assert patched.status_code == 422
        closed = (await client.patch(f"/engagements/{eng['id']}", json={"closed": True, "business_unit": "Retro Rabbit"})).json()
        assert closed["closed"] is True and closed["business_unit"] == "Retro Rabbit"
        view = (await client.get("/portfolio/forecast?month=2026-07")).json()
        assert eng["id"] not in {r["engagement_id"] for r in view["rows"]} and view["groups"] == []
        listed = (await client.get("/engagements")).json()["items"]
        assert next(e for e in listed if e["id"] == eng["id"])["closed"] is True


@pytest.mark.parametrize("role, can_write", [("Presales", False), ("HR", False), ("Delivery", True), ("Finance", True)])
async def test_permissions_on_the_new_endpoints(session, role, can_write):
    await seed_calendar(session)
    writer = await actor_for(session, ("Delivery",))
    client_row = await seed_client(session, "Client A")
    async with client_for(session, writer) as client:
        eng, allocations = await seed_client_a(client, client_row)
    actor = await actor_for(session, (role,))
    async with client_for(session, actor) as client:
        put = await client.put(f"/allocations/{allocations[0]['id']}/actuals", json=[])
        imp = await client.post("/engagements/import-forecast-model", files={"file": ("m.xlsx", model_bytes())}, data={"dry_run": "true"})
        assert (put.status_code == 200) is can_write and (imp.status_code == 200) is can_write, (role, put.status_code, imp.status_code)


async def test_model_import_rejects_a_workbook_that_is_not_the_master(session):
    await seed_calendar(session)
    actor = await actor_for(session, ("Delivery",))
    sheet = (Path(__file__).resolve().parents[2] / "docs/reference/DealGate_Cost_Margin_SAMPLE_StaffAug.xlsx").read_bytes()
    async with client_for(session, actor) as client:
        response = await client.post("/engagements/import-forecast-model", files={"file": ("sheet.xlsx", sheet)}, data={"dry_run": "true"})
        assert response.status_code == 422 and "Not a forecast model master" in response.text
        response = await client.post("/engagements/import-forecast-model", files={"file": ("x.xlsx", b"not a workbook")}, data={"dry_run": "true"})
        assert response.status_code == 422


async def test_as_of_defaults_to_the_models_month_and_never_leaks_from_an_earlier_request(session):
    """P5b: a request without as_of resolves to the importer's MODEL_AS_OF (2026-09) even right after a
    request that asked for a different month — the service passes None through; nothing is cached."""
    await seed_calendar(session)
    actor = await actor_for(session, ("Delivery",))
    async with client_for(session, actor) as client:
        earlier = await upload_model(client, dry_run=True, as_of="2025-10")
        assert earlier["report_markdown"].splitlines()[0].startswith("# ")
        assert "Actuals are taken through 2025-10" in earlier["report_markdown"]
        alpha_earlier = next(t for t in earlier["tabs"] if t["tab"] == "Alpha")
        default = await upload_model(client, dry_run=True, as_of=None)
        assert "Actuals are taken through 2026-09; every later month is forecast" in default["report_markdown"]
        alpha_default = next(t for t in default["tabs"] if t["tab"] == "Alpha")
        assert alpha_earlier["actual_months"] == 2 and alpha_default["actual_months"] == 3     # Sep+Oct vs Sep+Oct+Nov 2025
        bad = await client.post("/engagements/import-forecast-model", files={"file": ("m.xlsx", model_bytes())}, data={"dry_run": "true", "as_of": "Sept"})
        assert bad.status_code == 422


async def test_a_load_with_blockers_is_422_and_writes_nothing(session):
    """P5c: the fixture's 'Zeta Unlisted' tab has no dashboard row and a blank location. The dry run reports both
    (200, nothing created); the load refuses with 422 {blockers} before any write — no client, engagement,
    person, location or holiday appears; nothing is filed under a default business unit."""
    from app.models.engagement import CalendarLocation, Holiday, Resource

    await seed_calendar(session)
    actor = await actor_for(session, ("SystemAdmin",))
    async with client_for(session, actor) as client:
        async def counts():
            out = []
            for m in (Client, Engagement, Resource, CalendarLocation, Holiday):
                out.append(await session.scalar(select(func.count()).select_from(m)))
            return tuple(out)

        before = await counts()
        dry = await upload_model(client, dry_run=True)
        assert dry["dry_run"] is True and dry["created"] is None
        kinds = {(b["kind"], b["tab"]) for b in dry["blockers"]}
        assert kinds == {("business_unit", "Zeta Unlisted"), ("location", "Zeta Unlisted")}
        loc = next(b for b in dry["blockers"] if b["kind"] == "location")
        assert loc["rows"] and loc["fix"].startswith("Fill the Location column") and all(name in loc["fix"] for name in loc["rows"])
        bu = next(b for b in dry["blockers"] if b["kind"] == "business_unit")
        assert "Dashboard (RunRate)" in bu["fix"]
        assert dry["report_markdown"].index('> **This workbook ca') < dry["report_markdown"].index("Actuals are taken through")
        response = await client.post(
            "/engagements/import-forecast-model",
            files={"file": ("model.xlsx", model_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"dry_run": "false", "as_of": "2026-02"},
        )
        assert response.status_code == 422, response.text
        body = response.json()["detail"]
        assert "errors" not in body and {(b["kind"], b["tab"]) for b in body["blockers"]} == kinds
        assert await counts() == before
        assert (await client.get("/engagements?page_size=50")).json()["total"] == 0
        # the same workbook without that tab loads
        loaded = await upload_model(client, dry_run=False, complete=True)
        assert loaded["blockers"] == [] and loaded["created"]["engagements"] == 5
        assert all(e["business_unit"] in ("ST21", "IT Avalon", "Retro Rabbit") for e in (await client.get("/engagements?page_size=50")).json()["items"])


async def test_placeholder_seats_and_extensions_survive_a_second_import(session):
    """P5 final, item 3: the importer's row_key is the identity (allocation.import_key). Importing the fixed fixture
    twice leaves Zeta Unlisted with four allocations both times — Person Z1, Person Z2, TBC #1 and TBC #2 — nothing
    duplicated, nothing collapsed; and Gamma's Person G4, whose extension is typed as starting on the day the original
    ends, keeps a first segment that ends 2026-03-30 (the newer contract wins from its start)."""
    from app.services.forecast.forecast_model_import import parse_master

    await seed_calendar(session)
    actor = await actor_for(session, ("SystemAdmin",))
    assert parse_master(io.BytesIO(fixed_bytes()), as_of="2026-02").blockers == []
    async with client_for(session, actor) as client:
        first = await upload_model(client, dry_run=False, fixed=True)
        assert first["blockers"] == [] and first["created"]["engagements"] == 6
        listed = (await client.get("/engagements?page_size=50")).json()["items"]
        zeta = next(e for e in listed if e["name"].startswith("Zeta Unlisted"))
        assert zeta["business_unit"] == "ST21" and zeta["sales_owner"] == "Sam"

        async def zeta_rows():
            rows = (await client.get(f"/engagements/{zeta['id']}/allocations")).json()
            return sorted(r["import_key"] for r in rows), rows

        keys_1, rows_1 = await zeta_rows()
        assert len(rows_1) == 4 and all(k for k in keys_1) and len(set(keys_1)) == 4
        assert any(k.endswith(":TBC #1:0") for k in keys_1) and any(k.endswith(":TBC #2:0") for k in keys_1)
        assert {r["person"] for r in rows_1} == {"Person Z1", "Person Z2", "TBC"}
        total_before = await session.scalar(select(func.count()).select_from(Allocation))

        second = await upload_model(client, dry_run=False, fixed=True)
        assert second["created"]["engagements"] == 0 and second["created"]["allocations"] == 0 and second["created"]["clients"] == 0
        keys_2, rows_2 = await zeta_rows()
        assert keys_2 == keys_1 and len(rows_2) == 4
        assert {r["id"] for r in rows_2} == {r["id"] for r in rows_1}                       # the same rows, updated in place
        assert await session.scalar(select(func.count()).select_from(Allocation)) == total_before

        gamma_vms = next(e for e in listed if e["name"].startswith("Gamma Health-VMS") and "VMS" in e["name"].split("·")[-1])
        g4 = [r for r in (await client.get(f"/engagements/{gamma_vms['id']}/allocations")).json() if r["person"] == "Person G4"]
        assert sorted((r["import_key"].rsplit(":", 1)[-1], r["effective_start"], r["effective_end"]) for r in g4) == [
            ("0", "2025-12-01", "2026-03-30"),
            ("1", "2026-03-31", "2026-09-30"),
        ]
