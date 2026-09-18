"""S6 legacy-import: end-to-end vertical slice.

Acceptance tests mirrored from ``docs/backlog/s6-legacy-import.md`` and
``docs/backlog/s6-legacy-reconciliation.md``:

- Bulk upload creates sow_versions with legacy=true / approval_evidenced=false.
- Excel import rejects unknown location.
- Excel import rejects missing hourly_loaded_cost (§2 rule).
- Uploaded 5 SOWs vs Excel with only 3 matching → reconciliation lists 2 unmatched.
- Approve batch → transitions status + files follow-up tasks.
- Non-Finance/CEO/SystemAdmin → 403 on every endpoint.
- Attempting an approval on a legacy sow_version → 403 with the rollout rule
  message (via the placeholder guard).
"""

from __future__ import annotations

import io
import uuid
from pathlib import Path

import httpx
import pytest
import pytest_asyncio
from openpyxl import Workbook
from sqlalchemy import select

from app.db import get_session
from app.main import app as main_app
from app.models.gm_model import GmModel, ResourceLine
from app.models.legacy import LegacyImportBatch
from app.models.sow import Sow, SowVersion
from app.models.task import Task
from app.models.user import User
from app.services.legacy_import import (
    ALLOWED_LOCATIONS,
    LegacyImportError,
    SowUpload,
    assert_not_legacy_for_approval,
    bulk_upload_sows,
    create_batch,
    import_excel,
    reconcile,
    approve_batch,
)


FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "legacy_projects"


def _uid(email: str) -> uuid.UUID:
    return uuid.uuid5(uuid.NAMESPACE_URL, f"dealgate:local:{email}")


def _client(app):
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    )


@pytest.fixture(autouse=True)
def _local_env(monkeypatch):
    monkeypatch.setenv("DEALGATE_ENV", "local")
    monkeypatch.delenv("DEALGATE_TEST_GROUPS", raising=False)


@pytest_asyncio.fixture
async def app_with_session(session):
    async def _override():
        yield session

    main_app.dependency_overrides[get_session] = _override
    try:
        yield main_app
    finally:
        main_app.dependency_overrides.pop(get_session, None)


@pytest_asyncio.fixture
async def finance_user(session) -> User:
    u = User(
        id=_uid("finance@smartek21.com"),
        email="finance@smartek21.com",
        name="Finance User",
        groups=["Finance"],
    )
    session.add(u)
    await session.commit()
    await session.refresh(u)
    return u


# ---- helpers ------------------------------------------------------------


def _make_xlsx(rows: list[dict]) -> bytes:
    """Build an in-memory workbook mirroring the template columns."""

    columns = (
        "sow_ref",
        "client_name",
        "engagement_type",
        "role",
        "seniority",
        "location",
        "start_date",
        "end_date",
        "allocation_pct",
        "billable_hours",
        "hourly_bill_rate",
        "hourly_loaded_cost",
        "currency",
        "revenue_us",
        "revenue_india",
        "notes",
    )
    wb = Workbook()
    ws = wb.active
    for c, name in enumerate(columns, start=1):
        ws.cell(row=1, column=c, value=name)
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(columns, start=1):
            ws.cell(row=r, column=c, value=row.get(name))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _good_row(sow_ref: str = "SOW-X-1", **overrides) -> dict:
    base = {
        "sow_ref": sow_ref,
        "client_name": "Acme",
        "engagement_type": "staff_aug",
        "role": "Engineer",
        "seniority": "senior",
        "location": "US",
        "start_date": "2026-01-01",
        "end_date": "2026-03-31",
        "allocation_pct": 100,
        "billable_hours": 100,
        "hourly_bill_rate": 200,
        "hourly_loaded_cost": 120,
        "currency": "USD",
        "revenue_us": None,
        "revenue_india": None,
        "notes": None,
    }
    base.update(overrides)
    return base


# ---- role gating -------------------------------------------------------


async def test_non_finance_forbidden_on_all_endpoints(
    app_with_session, session, monkeypatch, finance_user
):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Sales")
    async with _client(app_with_session) as c:
        for method, path, body in [
            ("POST", "/legacy/upload-url", {"filename": "x.pdf", "content_type": "application/pdf"}),
            ("POST", "/legacy/batches", None),
            (
                "POST",
                f"/legacy/batches/{uuid.uuid4()}/sows",
                {"files": []},
            ),
            ("GET", f"/legacy/batches/{uuid.uuid4()}/reconciliation", None),
            ("POST", f"/legacy/batches/{uuid.uuid4()}/approve", None),
        ]:
            if method == "GET":
                r = await c.get(path, headers={"X-Test-User": "sales@x.com"})
            else:
                r = await c.post(
                    path,
                    headers={"X-Test-User": "sales@x.com"},
                    json=body,
                )
            assert r.status_code == 403, (method, path, r.text)


# ---- POST /legacy/batches + bulk_upload_sows ---------------------------


async def test_bulk_upload_creates_legacy_sow_versions(session, finance_user):
    batch = await create_batch(session, actor_id=finance_user.id)
    uploads = [
        SowUpload(
            s3_key=f"sow/x/2026/{i}.pdf",
            filename=f"sow-{i}.pdf",
            sow_ref=f"SOW-CLIENT-A-{i}",
            client_name="Acme",
        )
        for i in range(3)
    ]
    versions = await bulk_upload_sows(
        session, actor_id=finance_user.id, batch=batch, uploads=uploads
    )
    assert len(versions) == 3
    for v in versions:
        # governance flags
        assert v.legacy is True  # type: ignore[attr-defined]
        assert v.approval_evidenced is False  # type: ignore[attr-defined]
    assert batch.sow_count == 3
    assert batch.status == "reviewing"


# ---- Excel import validation -------------------------------------------


async def test_excel_import_rejects_unknown_location(session, finance_user):
    batch = await create_batch(session, actor_id=finance_user.id)
    xlsx = _make_xlsx([_good_row(location="EU")])
    with pytest.raises(LegacyImportError) as ei:
        await import_excel(
            session, actor_id=finance_user.id, batch=batch, xlsx_bytes=xlsx
        )
    err_cols = {e.get("column") for e in ei.value.errors}
    assert "location" in err_cols
    # All-or-nothing: no ResourceLine rows exist.
    rows = (await session.execute(select(ResourceLine))).scalars().all()
    assert list(rows) == []


async def test_excel_import_rejects_missing_hourly_loaded_cost(session, finance_user):
    batch = await create_batch(session, actor_id=finance_user.id)
    xlsx = _make_xlsx([_good_row(hourly_loaded_cost=None)])
    with pytest.raises(LegacyImportError) as ei:
        await import_excel(
            session, actor_id=finance_user.id, batch=batch, xlsx_bytes=xlsx
        )
    err_cols = {e.get("column") for e in ei.value.errors}
    assert "hourly_loaded_cost" in err_cols
    rows = (await session.execute(select(ResourceLine))).scalars().all()
    assert list(rows) == []


async def test_excel_import_happy_path_creates_gm_models(session, finance_user):
    batch = await create_batch(session, actor_id=finance_user.id)
    # Pre-attach a matching SOW so the model links up.
    await bulk_upload_sows(
        session,
        actor_id=finance_user.id,
        batch=batch,
        uploads=[
            SowUpload(
                s3_key="sow/x/1.pdf",
                filename="1.pdf",
                sow_ref="SOW-X-1",
                client_name="Acme",
            )
        ],
    )
    xlsx = _make_xlsx([_good_row(), _good_row(role="QA", location="India", hourly_bill_rate=50, hourly_loaded_cost=20, billable_hours=50)])
    result = await import_excel(
        session, actor_id=finance_user.id, batch=batch, xlsx_bytes=xlsx
    )
    assert result["imported"] == 2
    models = (await session.execute(select(GmModel))).scalars().all()
    assert len(models) == 1
    lines = (await session.execute(select(ResourceLine))).scalars().all()
    assert len(lines) == 2


# ---- Reconciliation ----------------------------------------------------


async def test_reconciliation_reports_unmatched_sows(session, finance_user):
    batch = await create_batch(session, actor_id=finance_user.id)
    # Upload 5 SOWs
    uploads = [
        SowUpload(
            s3_key=f"sow/x/{i}.pdf",
            filename=f"sow-{i}.pdf",
            sow_ref=f"SOW-CLIENT-A-{i}",
            client_name="Acme",
        )
        for i in range(5)
    ]
    await bulk_upload_sows(
        session, actor_id=finance_user.id, batch=batch, uploads=uploads
    )
    # Excel with only 3 matching sow_refs.
    rows = [
        _good_row(sow_ref=f"SOW-CLIENT-A-{i}") for i in range(3)
    ]
    await import_excel(session, actor_id=finance_user.id, batch=batch, xlsx_bytes=_make_xlsx(rows))
    report = await reconcile(session, batch.id)
    assert len(report.matched) == 3
    assert len(report.unmatched_sows) == 2


# ---- Approve + tasks ---------------------------------------------------


async def test_approve_batch_creates_tasks_for_unmatched(session, finance_user):
    batch = await create_batch(session, actor_id=finance_user.id)
    await bulk_upload_sows(
        session,
        actor_id=finance_user.id,
        batch=batch,
        uploads=[
            SowUpload(
                s3_key=f"sow/x/{i}.pdf",
                filename=f"sow-{i}.pdf",
                sow_ref=f"SOW-CLIENT-A-{i}",
                client_name="Acme",
            )
            for i in range(2)
        ],
    )
    await import_excel(
        session,
        actor_id=finance_user.id,
        batch=batch,
        xlsx_bytes=_make_xlsx([_good_row(sow_ref="SOW-CLIENT-A-0")]),
    )
    before_tasks = (await session.execute(select(Task))).scalars().all()
    report = await approve_batch(session, actor_id=finance_user.id, batch_id=batch.id)
    assert report.status == "approved"
    after_tasks = (await session.execute(select(Task))).scalars().all()
    assert len(after_tasks) > len(before_tasks)


# ---- Legacy approval guard --------------------------------------------


async def test_assert_not_legacy_for_approval_rejects_legacy_rows(session, finance_user):
    from fastapi import HTTPException

    batch = await create_batch(session, actor_id=finance_user.id)
    versions = await bulk_upload_sows(
        session,
        actor_id=finance_user.id,
        batch=batch,
        uploads=[
            SowUpload(
                s3_key="sow/y/1.pdf",
                filename="y.pdf",
                sow_ref="SOW-Y-1",
                client_name="Beta",
            )
        ],
    )
    with pytest.raises(HTTPException) as ei:
        assert_not_legacy_for_approval(versions[0])
    assert ei.value.status_code == 403
    assert "rollout" in ei.value.detail


# ---- Router flow (create batch + attach + reconciliation) --------------


async def test_router_create_batch_returns_row(app_with_session, session, monkeypatch, finance_user):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Finance")
    async with _client(app_with_session) as c:
        r = await c.post("/legacy/batches", headers={"X-Test-User": "finance@smartek21.com"})
    assert r.status_code == 201
    body = r.json()
    assert body["status"] == "uploading"
    assert body["sow_count"] == 0


async def test_router_upload_url_returns_signed_url(
    app_with_session, session, monkeypatch, finance_user
):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Finance")
    async with _client(app_with_session) as c:
        r = await c.post(
            "/legacy/upload-url",
            headers={"X-Test-User": "finance@smartek21.com"},
            json={"filename": "acme.pdf", "content_type": "application/pdf"},
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["s3_key"].endswith(".pdf")
    assert body["method"] == "PUT"


async def test_router_excel_import_422_on_bad_location(
    app_with_session, session, monkeypatch, finance_user
):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Finance")
    async with _client(app_with_session) as c:
        r = await c.post("/legacy/batches", headers={"X-Test-User": "finance@smartek21.com"})
        batch_id = r.json()["id"]
        xlsx = _make_xlsx([_good_row(location="Nowhere")])
        r = await c.post(
            f"/legacy/batches/{batch_id}/excel",
            headers={"X-Test-User": "finance@smartek21.com"},
            files={"file": ("bad.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert r.status_code == 422
    body = r.json()
    assert "location" in str(body["detail"])


async def test_router_reconciliation_returns_report(
    app_with_session, session, monkeypatch, finance_user
):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Finance")
    async with _client(app_with_session) as c:
        r = await c.post("/legacy/batches", headers={"X-Test-User": "finance@smartek21.com"})
        batch_id = r.json()["id"]
        # attach 2 sows
        r = await c.post(
            f"/legacy/batches/{batch_id}/sows",
            headers={"X-Test-User": "finance@smartek21.com"},
            json={
                "files": [
                    {"s3_key": "sow/x/1.pdf", "filename": "1.pdf", "sow_ref": "S-1", "client_name": "Acme"},
                    {"s3_key": "sow/x/2.pdf", "filename": "2.pdf", "sow_ref": "S-2", "client_name": "Acme"},
                ]
            },
        )
        assert r.status_code == 200, r.text
        # excel with one matching row
        xlsx = _make_xlsx([_good_row(sow_ref="S-1")])
        r = await c.post(
            f"/legacy/batches/{batch_id}/excel",
            headers={"X-Test-User": "finance@smartek21.com"},
            files={
                "file": (
                    "sample.xlsx",
                    xlsx,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert r.status_code == 200, r.text
        r = await c.get(
            f"/legacy/batches/{batch_id}/reconciliation",
            headers={"X-Test-User": "finance@smartek21.com"},
        )
    assert r.status_code == 200
    body = r.json()
    assert body["status"] in ("uploading", "reviewing")
    assert len(body["matched"]) == 1
    assert len(body["unmatched_sows"]) == 1
