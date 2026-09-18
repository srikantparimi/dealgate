"""GM sandbox API — S2 E4 (M1 sign-off screen).

This is the automated M1 evidence gate: Finance signs off on the six-template
math via the sandbox before Sales sees any GM UI. The §7 discounted case
(US $90k/$65k, India $55k/$30k) is asserted end-to-end via HTTP.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import httpx
import openpyxl
import pytest

from app.main import app as main_app


@pytest.fixture(autouse=True)
def _local_env(monkeypatch):
    monkeypatch.setenv("DEALGATE_ENV", "local")
    monkeypatch.delenv("DEALGATE_TEST_GROUPS", raising=False)


def _client(app):
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")


def _as_finance(monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Finance")


def _as_sales(monkeypatch):
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "Sales")


def _discounted_case_payload() -> dict:
    """§7 mixed-geography discounted case: US 90k rev / 65k cost,
    India 55k rev / 30k cost. Both components fail their floors."""
    return {
        "engagement_type": "fixed_price",
        "inputs": {
            "total_price": "145000",
            "revenue_us": "90000",
            "revenue_india": "55000",
            "resources": [
                {
                    "role": "Engineer",
                    "seniority": "Sr",
                    "location": "US",
                    "allocation_pct": "1",
                    "start": "2026-01-01",
                    "end": "2026-06-30",
                    "hours_billable": "1000",
                    "hourly_bill_rate": "0",
                    "hourly_cost": "65",
                },
                {
                    "role": "Engineer",
                    "seniority": "Sr",
                    "location": "India",
                    "allocation_pct": "1",
                    "start": "2026-01-01",
                    "end": "2026-06-30",
                    "hours_billable": "600",
                    "hourly_bill_rate": "0",
                    "hourly_cost": "50",
                },
            ],
        },
    }


# --- role gate -------------------------------------------------------------


async def test_sandbox_requires_auth():
    async with _client(main_app) as c:
        r = await c.post("/gm/sandbox", json=_discounted_case_payload())
    assert r.status_code == 401


async def test_sandbox_forbidden_for_sales(monkeypatch):
    _as_sales(monkeypatch)
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "sales@smartek21.com"},
            json=_discounted_case_payload(),
        )
    assert r.status_code == 403


async def test_sandbox_forbidden_for_ceo(monkeypatch):
    """CEO sees dashboards, not the sandbox math tool."""
    monkeypatch.setenv("DEALGATE_TEST_GROUPS", "CEO")
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "ceo@smartek21.com"},
            json=_discounted_case_payload(),
        )
    assert r.status_code == 403


async def test_sandbox_allowed_for_finance(monkeypatch):
    _as_finance(monkeypatch)
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "finance@smartek21.com"},
            json=_discounted_case_payload(),
        )
    assert r.status_code == 200


async def test_sandbox_allowed_for_delivery_presales_sysadmin(monkeypatch):
    for role in ("Delivery", "Presales", "SystemAdmin"):
        monkeypatch.setenv("DEALGATE_TEST_GROUPS", role)
        async with _client(main_app) as c:
            r = await c.post(
                "/gm/sandbox",
                headers={"X-Test-User": f"{role.lower()}@smartek21.com"},
                json=_discounted_case_payload(),
            )
        assert r.status_code == 200, role


# --- §7 discounted case: the M1 automated evidence -------------------------


async def test_sandbox_section7_discounted_case_is_m1_evidence(monkeypatch):
    """M1 automated evidence: §7 discounted case returns US GM 27.78%,
    India GM 45.45%, both fail floor, requires_ceo=True."""
    _as_finance(monkeypatch)
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "finance@smartek21.com"},
            json=_discounted_case_payload(),
        )
    assert r.status_code == 200
    body = r.json()

    # Component revenue/cost round-trip
    assert Decimal(body["revenue_us"]) == Decimal("90000")
    assert Decimal(body["cost_us"]) == Decimal("65000")
    assert Decimal(body["revenue_india"]) == Decimal("55000")
    assert Decimal(body["cost_india"]) == Decimal("30000")

    # GM components — quantized to 4dp for display.
    assert Decimal(body["gm_us"]).quantize(Decimal("0.0001")) == Decimal("0.2778")
    assert Decimal(body["gm_india"]).quantize(Decimal("0.0001")) == Decimal("0.4545")

    # Both fail floor, requires CEO.
    assert body["complete"] is True
    assert body["policy"]["us_pass"] is False
    assert body["policy"]["india_pass"] is False
    assert body["policy"]["requires_ceo"] is True
    assert set(body["policy"]["failing"]) == {"US", "India"}


# --- missing input handling ------------------------------------------------


async def test_missing_hourly_cost_marks_incomplete(monkeypatch):
    _as_finance(monkeypatch)
    payload = _discounted_case_payload()
    # Drop hourly_cost on the first resource.
    payload["inputs"]["resources"][0].pop("hourly_cost")
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "finance@smartek21.com"},
            json=payload,
        )
    assert r.status_code == 200
    body = r.json()
    assert body["complete"] is False
    assert "resources[0].hourly_cost" in body["missing"]
    # GM on the incomplete side is null — never a spurious zero.
    assert body["gm_us"] is None
    # India side still computes.
    assert Decimal(body["gm_india"]).quantize(Decimal("0.0001")) == Decimal("0.4545")
    # Incomplete sheets require CEO.
    assert body["policy"]["requires_ceo"] is True


async def test_invalid_engagement_type_rejected(monkeypatch):
    _as_finance(monkeypatch)
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "finance@smartek21.com"},
            json={"engagement_type": "not_a_thing", "inputs": {}},
        )
    assert r.status_code == 422


# --- schema endpoint -------------------------------------------------------


async def test_schema_endpoint_returns_template_fields(monkeypatch):
    _as_finance(monkeypatch)
    async with _client(main_app) as c:
        r = await c.get(
            "/gm/sandbox/schema/fixed_price",
            headers={"X-Test-User": "finance@smartek21.com"},
        )
    assert r.status_code == 200
    schema = r.json()
    assert schema["engagement_type"] == "fixed_price"
    field_names = {f["name"] for f in schema["fields"]}
    assert {"total_price", "revenue_us", "revenue_india", "resources"} <= field_names


async def test_schema_endpoint_forbidden_for_sales(monkeypatch):
    _as_sales(monkeypatch)
    async with _client(main_app) as c:
        r = await c.get(
            "/gm/sandbox/schema/fixed_price",
            headers={"X-Test-User": "sales@smartek21.com"},
        )
    assert r.status_code == 403


async def test_schema_endpoint_rejects_unknown_type(monkeypatch):
    _as_finance(monkeypatch)
    async with _client(main_app) as c:
        r = await c.get(
            "/gm/sandbox/schema/nope",
            headers={"X-Test-User": "finance@smartek21.com"},
        )
    assert r.status_code == 422


# --- xlsx export -----------------------------------------------------------


async def test_export_xlsx_numbers_match_compute(monkeypatch):
    _as_finance(monkeypatch)
    payload = _discounted_case_payload()
    async with _client(main_app) as c:
        compute = await c.post(
            "/gm/sandbox",
            headers={"X-Test-User": "finance@smartek21.com"},
            json=payload,
        )
        export = await c.post(
            "/gm/sandbox/export",
            headers={"X-Test-User": "finance@smartek21.com"},
            json=payload,
        )
    assert compute.status_code == 200
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(BytesIO(export.content), data_only=False)
    assert set(wb.sheetnames) == {"Inputs", "Result"}
    result_sheet = wb["Result"]

    # Build a {label: value} dict from the "Result" sheet (col A = label, col B = value).
    kv: dict[str, object] = {}
    for row in result_sheet.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        kv[str(row[0])] = row[1]

    body = compute.json()
    assert Decimal(str(kv["revenue_us"])) == Decimal(body["revenue_us"])
    assert Decimal(str(kv["cost_us"])) == Decimal(body["cost_us"])
    assert Decimal(str(kv["revenue_india"])) == Decimal(body["revenue_india"])
    assert Decimal(str(kv["cost_india"])) == Decimal(body["cost_india"])
    assert Decimal(str(kv["gm_us"])) == Decimal(body["gm_us"])
    assert Decimal(str(kv["gm_india"])) == Decimal(body["gm_india"])
    # Policy pass/fail must match.
    assert str(kv["us_pass"]).lower() == str(body["policy"]["us_pass"]).lower()
    assert str(kv["india_pass"]).lower() == str(body["policy"]["india_pass"]).lower()
    assert str(kv["requires_ceo"]).lower() == str(body["policy"]["requires_ceo"]).lower()


async def test_export_forbidden_for_sales(monkeypatch):
    _as_sales(monkeypatch)
    async with _client(main_app) as c:
        r = await c.post(
            "/gm/sandbox/export",
            headers={"X-Test-User": "sales@smartek21.com"},
            json=_discounted_case_payload(),
        )
    assert r.status_code == 403
