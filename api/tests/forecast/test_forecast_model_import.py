"""The forecast-model master workbook (one tab per client, month blocks, several header layouts) imports into
engagements, allocations and actuals; the engine reproduces the sheet's own plan numbers.
The fixture is synthetic — same layouts as the real model, invented people and clients."""
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path
import pytest
from openpyxl import Workbook

from app.services.forecast import Calendar, compute
from app.services.forecast.forecast_model_import import parse_master, to_engine, dry_run_markdown, reconcile

FIX = Path(__file__).parent / "fixtures" / "forecast_model_synthetic.xlsx"


def build_fixture(path: Path):
    wb = Workbook(); wb.remove(wb.active)
    # --- dashboard ----------------------------------------------------------------------------------------
    ws = wb.create_sheet("Dashboard (RunRate)")
    ws["B5"], ws["C5"], ws["D5"], ws["E5"] = "Client", "Sales Owner", "Delivery Owner", "Billing Type"
    rows = [("ST21",), ("Alpha Retail", "Sam", "Dee", "T&M"), ("Beta Insurance", "Sam", "Raj", "Fixed Monthly"), ("Subtotal",),
            ("IT Avalon",), ("Gamma Health-VMS", "Kim", "Lee", "T&M"), ("Delta Closed", "Kim", "Lee", "Closed"), ("Subtotal",),
            ("Retro Rabbit",), ("Epsilon", "Jo", "Jo", "T&M"), ("Subtotal",), ("Total",),
            (None,), ("Client",), ("ST21",), ("ITA",), ("RR",), ("Total",)]   # the roll-up block the real sheet repeats below the total
    for i, r in enumerate(rows, start=6):
        for j, v in enumerate(r, start=2): ws.cell(row=i, column=j, value=v)
    # --- working days / holidays -----------------------------------------------------------------------------
    ws = wb.create_sheet("WorkingDays")
    ws["I2"] = "HOLIDAYS"; ws["I3"], ws["L3"], ws["O3"], ws["R3"] = "IN - PUNE", "IN - CHENNAI", "US", "SA"; ws["V2"] = "Location List"; ws["V3"] = "IN PUNE"; ws["V4"] = "US"
    for r, d in enumerate([date(2025, 10, 2), date(2025, 10, 21), date(2026, 1, 1), date(2026, 1, 26)], start=4): ws.cell(row=r, column=9, value=datetime.combine(d, datetime.min.time()))
    for r, d in enumerate([date(2025, 10, 1), date(2025, 10, 2), date(2025, 10, 20), date(2026, 1, 1), date(2026, 1, 15)], start=4): ws.cell(row=r, column=12, value=datetime.combine(d, datetime.min.time()))
    for r, d in enumerate([date(2025, 11, 27), date(2025, 12, 25), date(2026, 1, 1), date(2026, 9, 7)], start=4): ws.cell(row=r, column=15, value=datetime.combine(d, datetime.min.time()))
    ws.cell(row=4, column=18, value=datetime(2025, 12, 16))

    def month_blocks(ws, first_col, months, kind_by_month, headers, year_in_row1=None):
        c = first_col
        if year_in_row1: ws.cell(row=1, column=c, value=year_in_row1)
        for m in months:
            ws.cell(row=2, column=c, value=kind_by_month(m)); ws.cell(row=3, column=c, value=m)
            for k, h in enumerate(headers): ws.cell(row=4, column=c + k, value=h)
            c += len(headers)
        return c

    # --- Alpha: Costco-style. Hours block then Days blocks, SOW dates in the Team/SOW text, Onsite/Offshore, flat max billing
    ws = wb.create_sheet("Alpha")
    ws["A2"] = "ALPHA RETAIL"
    for k, h in enumerate(["Resource", "Role", "Team/SOW", "Type", "SOW Start Date", "SOW End Date", "Location", "Rate"], start=1): ws.cell(row=4, column=k, value=h)
    months = ["September", "October", "November", "December", "January'26", "February'26", "March'26"]
    month_blocks(ws, 9, months[:1], lambda m: "Actual", ["Hours", "Vacation Hours/Planned Leaves", "Unplanned PTO Leave", "Max Billing", "Billing"], 2025)
    month_blocks(ws, 14, months[1:], lambda m: "Actual" if m in ("October", "November") else "Forecast", ["Days (Not Inclusive of Holidays)", "Vacation Days/Planned Leaves", "Unplanned PTO Leave", "Max Billing", "Billing"])
    sow = "Alpha Retail - SOW #12 - Professional Services - Travel Team 1 - 1-Sep-25 to 31-Aug-27"
    people = [("Person A1", "OC", "Onsite", 120), ("Person A2", "Dev", "Offshore", 32), ("Person A3", "Dev", "Offshore", 32)]
    for i, (n, role, loc, rate) in enumerate(people, start=5):
        for k, v in enumerate([n, role, sow, "Managed Services", None, None, loc, rate], start=1): ws.cell(row=i, column=k, value=v)
        ws.cell(row=i, column=9, value=173); ws.cell(row=i, column=10, value=8 if i == 6 else 0); ws.cell(row=i, column=12, value=173 * rate); ws.cell(row=i, column=13, value=(173 - (8 if i == 6 else 0)) * rate)
        c = 14
        for m, days in zip(months[1:], [21, 20, 22, 20, 20, 22]):
            ws.cell(row=i, column=c, value=days); ws.cell(row=i, column=c + 3, value=173 * rate)
            if m in ("October", "November"): ws.cell(row=i, column=c + 1, value=1 if i == 5 else None); ws.cell(row=i, column=c + 4, value=173 * rate * (days - (1 if i == 5 else 0)) / days)
            c += 5
    ws.cell(row=9, column=1, value="Total")

    # --- Beta: Liberty-style. Monthly rates per person, Allocation %, Days, no row dates (SOW term in the text)
    ws = wb.create_sheet("Beta")
    ws["A2"] = "BETA INSURANCE"
    for k, h in enumerate(["Resource", "Role", "Team/SOW", "Type", "Start Date", "End Date", "Location", "Rate", "Allocation %"], start=1): ws.cell(row=4, column=k, value=h)
    months = ["September", "October", "November", "December", "January'26"]
    month_blocks(ws, 10, months, lambda m: "Actual" if m != "January'26" else "Forecast", ["Days", "Vacation Days", "Unplanned Leave", "Max Billing", "Billing"], 2025)
    sow = "Beta Services for O365 Support - CR #01 - BT00002382 - 1-Jul-25 to 30-Jun-26"
    for i, (n, loc, rate, pct) in enumerate([("Person B1", "IN CHENNAI", 5151, 1), ("Person B2", "IN PUNE", 6325, 0.5)], start=5):
        for k, v in enumerate([n, "Developer", sow, "Managed Services", None, None, loc, rate, pct], start=1): ws.cell(row=i, column=k, value=v)
        c = 10
        for m, days in zip(months, [22, 20, 20, 22, 19]):
            ws.cell(row=i, column=c, value=days); ws.cell(row=i, column=c + 3, value=rate * pct)
            if m != "January'26": ws.cell(row=i, column=c + 4, value=rate * pct)
            c += 5
    ws.cell(row=7, column=1, value="Total")

    # --- Gamma-VMS: staffing. Original SOW + extension pairs, Commission, Budgeted/Actual Hours, labels without a year, an OVERTIME row
    ws = wb.create_sheet("Gamma Health-VMS")
    ws["A2"] = "IT Staffing"
    for k, h in enumerate(["Resource", "Role", "Team/SOW", "Type", "Original SOW Start Date", "Original SOW End Date", "Extension Start Date", "Extension End Date", "Location", "Rate", "Commission"], start=1): ws.cell(row=4, column=k, value=h)
    months = ["January", "February", "March", "April"]
    month_blocks(ws, 12, months, lambda m: "Forecast", ["Budgeted Hours", "Actual Hours", "Budget vs Actual", "Max Billing", "Actual Billing"])
    rows = [("Person G1", "Epic Analyst", "VMS", "Staffing", datetime(2025, 11, 3), datetime(2026, 2, 27), datetime(2026, 2, 28), datetime(2026, 8, 31), "US", 145, "Y"),
            ("Person G1-OVERTIME", "Epic Analyst", "VMS", "Staffing", None, None, None, None, None, 217.5, "N"),
            ("Person G2", "PM", "PO 26-00389", "Staffing", datetime(2026, 1, 12), datetime(2026, 12, 31), None, None, "US", 100, "N"),
            ("Person G3", "Analyst", "VMS", "Staffing", datetime(2025, 12, 1), datetime(2026, 3, 31), None, datetime(2026, 6, 30), "US", 90, "N"),   # extension with no start date
            ("Person G4", "Analyst", "VMS", "Staffing", datetime(2025, 12, 1), datetime(2026, 3, 31), datetime(2026, 3, 31), datetime(2026, 9, 30), "US", 90, "N")]   # extension typed as starting on the day the original ends
    for i, r in enumerate(rows, start=5):
        for k, v in enumerate(r, start=1): ws.cell(row=i, column=k, value=v)
    # Jan+Feb have actuals (actual hours + actual billing); Mar/Apr are plan only
    for i, rate in ((5, 145), (6, 217.5), (7, 100), (8, 90), (9, 90)):
        c = 12
        for m, (bud, act) in zip(months, [(168, 160), (160, 152), (176, None), (176, None)]):
            if i == 6: bud, act = None, (4 if m == "January" else None)
            ws.cell(row=i, column=c, value=bud); ws.cell(row=i, column=c + 1, value=act)
            if bud: ws.cell(row=i, column=c + 3, value=bud * rate)
            if act is not None: ws.cell(row=i, column=c + 4, value=act * rate)
            c += 5
    ws.cell(row=10, column=1, value="Total")
    # --- Zeta Unlisted: not on the dashboard at all, and one person with no Location
    ws = wb.create_sheet("Zeta Unlisted"); ws["A2"] = "ZETA"
    for k, h in enumerate(["Resource", "Role", "Team/SOW", "Type", "SOW Start Date", "SOW End Date", "Location", "Rate", "Max Days", "Vacation Days", "Unplanned Leave", "Max Billing", "Billing"], start=1): ws.cell(row=4, column=k, value=h)
    ws["I1"], ws["I2"], ws["I3"] = 2026, "Actual", "January"
    for k, v in enumerate(["Person Z1", "Engineer", "SOW #4", "Managed Services", datetime(2026, 1, 1), datetime(2026, 6, 30), None, 60, 20, None, None, 9600, 9600], start=1): ws.cell(row=5, column=k, value=v)
    for k, v in enumerate(["Person Z2", "Engineer", "SOW #4", "Managed Services", datetime(2026, 1, 1), datetime(2026, 6, 30), "US", 60, 20, None, None, 9600, 9600], start=1): ws.cell(row=6, column=k, value=v)
    # two placeholder seats with the same name, and an invoice (8,400 = 140 h) that is not units-minus-leave (20 d - 1 d = 152 h)
    for k, v in enumerate(["TBC", "Engineer", "SOW #4", "Managed Services", datetime(2026, 1, 1), datetime(2026, 6, 30), "US", 60, 20, 1, None, 9600, 8400], start=1): ws.cell(row=7, column=k, value=v)
    for k, v in enumerate(["TBC", "Engineer", "SOW #4", "Managed Services", datetime(2026, 1, 1), datetime(2026, 6, 30), "US", 60, 20, None, None, 9600, 9600], start=1): ws.cell(row=8, column=k, value=v)
    ws["A9"] = "Total"

    # --- Delta Closed: CLOSED in the title
    ws = wb.create_sheet("Delta Closed"); ws["A2"] = "DELTA - CLOSED"
    for k, h in enumerate(["Resource", "Role", "Team/SOW", "Type", "SOW Start Date", "SOW End Date", "Location", "Rate", "Max Days", "Vacation Days", "Unplanned Leave", "Max Billing", "Billing"], start=1): ws.cell(row=4, column=k, value=h)
    ws["I2"], ws["I3"] = "Actual", "September"; ws["I1"] = 2025
    for k, v in enumerate(["Person D1", "Engineer", "Infra", "Managed Services", datetime(2025, 9, 1), datetime(2025, 9, 30), "SA", 28, 22, None, None, 4928, 4928], start=1): ws.cell(row=5, column=k, value=v)
    ws["A6"] = "Total"
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)


@pytest.fixture(scope="module")
def imp():
    build_fixture(FIX)
    return parse_master(FIX, as_of="2026-02")

@pytest.fixture(scope="module")
def cal(imp):
    return Calendar(D("8"), D("0.35"), {("US", "w2_employee"): D("0.36")}, {}, {k: set(v) for k, v in imp.locations.items()})


def test_the_roll_up_block_below_the_total_is_not_a_client_list(imp):
    """The real dashboard repeats Client / ST21 / ITA / RR as a summary under the grand total.
    Counting those as ST21 clients made the import screen read 9 ST21 clients instead of 7."""
    assert set(imp.dashboard) == {"Alpha Retail", "Beta Insurance", "Gamma Health-VMS", "Delta Closed", "Epsilon"}
    assert not any(k in imp.dashboard for k in ("ITA", "RR", "Total", "Client"))
    assert sum(1 for v in imp.dashboard.values() if v["business_unit"] == "ST21") == 2


def test_a_workbook_that_does_not_say_where_people_belong_blocks_the_load(imp):
    """A blank business unit or a blank location does not fail loudly once loaded — the rows land under a
    default, the totals still reconcile, and only the margin underneath is wrong. So the import refuses and
    names the cells. Nothing here is a judgement call; each one is an empty cell in the workbook."""
    kinds = {(b.kind, b.tab) for b in imp.blockers}
    assert ("business_unit", "Zeta Unlisted") in kinds          # no dashboard row -> no BU, no owners
    assert ("location", "Zeta Unlisted") in kinds
    loc = next(b for b in imp.blockers if b.kind == "location" and b.tab == "Zeta Unlisted")
    assert loc.rows == ["Person Z1"] and "Person Z2" not in loc.rows
    assert "1 person has no Location" in loc.detail and "US burden" in loc.detail
    assert "Fill the Location column" in loc.fix and "Person Z1" in loc.fix
    bu = next(b for b in imp.blockers if b.kind == "business_unit" and b.tab == "Zeta Unlisted")
    assert "Dashboard (RunRate)" in bu.fix
    # the tabs the workbook does describe are not blocked
    assert not any(b.tab in ("Alpha", "Beta") for b in imp.blockers)


def test_a_complete_workbook_has_no_blockers(tmp_path):
    """The same fixture without the unlisted tab loads with nothing to fix — the refusal is about missing
    facts, not a tax on every import."""
    import openpyxl
    path = tmp_path / "complete.xlsx"
    build_fixture(path)
    wb = openpyxl.load_workbook(path); wb.remove(wb["Zeta Unlisted"]); wb.save(path)
    assert parse_master(path, as_of="2026-02").blockers == []


def test_the_report_leads_with_the_blockers(imp):
    md = dry_run_markdown(imp)
    assert "This workbook cannot be loaded yet" in md
    assert md.index("cannot be loaded yet") < md.index("Actuals are taken through")
    assert "Zeta Unlisted" in md


def test_rows_that_share_a_name_keep_separate_identities(imp, cal):
    """Costco carries eleven 'TBC' seats on one SOW. A re-import keyed on the name would collapse them into one
    seat and drop ten people's forecast; keyed on row_key it finds each seat again."""
    z = next(e for e in imp.engagements if e.tab == "Zeta Unlisted")
    assert [a.row_key for a in z.allocations] == ["Person Z1", "Person Z2", "TBC #1", "TBC #2"]
    eng, allocs, acts, _ = to_engine(z, cal)
    assert len({a.id for a in allocs}) == 4 and {a.person for a in allocs} == {"Person Z1", "Person Z2", "TBC"}


def test_actual_hours_come_from_the_invoice_not_the_leave_columns(imp, cal):
    """The model's leave columns say 'hours' and hold days, or the reverse, and Costco's Days column is the
    calendar rather than what was billed: 1,606 of 2,705 real actual months disagreed with the invoice by that
    route. Where the sheet has an hourly rate and a billed amount, hours = billed / rate."""
    z = next(e for e in imp.engagements if e.tab == "Zeta Unlisted")
    eng, allocs, acts, _ = to_engine(z, cal)
    tbc1 = next(a for a in allocs if a.id.endswith("TBC #1:0"))
    jan = next(x for x in acts if x.allocation_id == tbc1.id and x.month == "2026-01")
    assert jan.hours == D("140") and jan.revenue == D("8400")          # not (20 - 1) × 8 = 152


def test_an_extension_with_no_start_begins_the_day_after_the_previous_segment(imp):
    g3 = [a for e in imp.engagements if e.tab == "Gamma Health-VMS" for a in e.allocations if a.resource == "Person G3"]
    assert [(a.segment, a.start, a.end) for a in g3] == [(0, date(2025, 12, 1), date(2026, 3, 31)), (1, date(2026, 4, 1), date(2026, 6, 30))]
    assert g3[1].row_key == g3[0].row_key == "Person G3"               # one identity across its segments


def test_an_extension_typed_to_start_on_the_original_s_last_day_does_not_bill_that_day_twice(imp):
    """Nine of the ten overlapping extensions in the real model are this: start == previous end. The newer
    contract wins from its start; the original gives way the day before. No typed date is altered."""
    g4 = [a for e in imp.engagements if e.tab == "Gamma Health-VMS" for a in e.allocations if a.resource == "Person G4"]
    assert [(a.segment, a.start, a.end) for a in g4] == [(0, date(2025, 12, 1), date(2026, 3, 30)), (1, date(2026, 3, 31), date(2026, 9, 30))]
    assert not g4[1].notes                                             # a one-day overlap is the sheet's convention, not an issue


def test_the_report_names_what_it_assumed(imp):
    md = dry_run_markdown(imp)
    assert "## Assumed, because the workbook does not say" in md
    for s in ("Employment type", "Offshore", "Cost per hour", "Currency", "Actual hours"):
        assert s in md


def test_a_sheet_is_read_once_and_only_if_it_is_needed():
    """The workbook is streamed, not parsed whole: openpyxl's normal mode cost ~5 s locally for the real
    model (~25 s on a staging task), nearly all of an import request. Sheets materialise on first touch."""
    from app.services.forecast.forecast_model_import import _Book
    bk = _Book(FIX)
    assert bk._grids == {}                                  # opening reads nothing
    alpha = bk["Alpha"]
    assert set(bk._grids) == {"Alpha"}
    assert bk["Alpha"] is alpha                             # and only once
    assert alpha.cell(row=4, column=1).value == "Resource"
    assert alpha.cell(row=4, column=999).value is None      # off the edge is blank, not an error
    assert alpha.cell(row=99999, column=1).value is None
    bk.close()


def test_as_of_is_per_call_and_never_leaks_between_imports():
    """One request asking for a different as-of must not change what the next request means by 'actual'.
    (It used to: as_of lived in a module global that every parse rewrote, so a worker inherited the last caller's.)"""
    from app.services.forecast import forecast_model_import as fm
    def actual_months(i):
        return sorted({m.month for e in i.engagements for a in e.allocations for m in a.months if m.kind == "actual"})
    earlier = fm.parse_master(FIX, as_of="2025-10")         # a caller asking for an older cut
    default_after = fm.parse_master(FIX)                    # the next caller, taking the default
    assert earlier.as_of == "2025-10" and actual_months(earlier) == ["2025-09", "2025-10"]
    assert default_after.as_of == fm.MODEL_AS_OF == "2026-09"
    assert actual_months(default_after) == ["2025-09", "2025-10", "2025-11", "2025-12", "2026-01", "2026-02"]
    assert actual_months(default_after) == actual_months(fm.parse_master(FIX, as_of="2026-09"))


def test_the_report_says_which_month_the_actuals_run_to(imp):
    assert f"Actuals are taken through {imp.as_of}" in dry_run_markdown(imp)


def test_dashboard_gives_bu_owners_and_billing_type(imp):
    assert imp.dashboard["Alpha Retail"] == dict(business_unit="ST21", sales_owner="Sam", delivery_owner="Dee", billing_type="T&M")
    assert imp.dashboard["Gamma Health-VMS"]["business_unit"] == "IT Avalon" and imp.dashboard["Epsilon"]["business_unit"] == "Retro Rabbit"
    assert any("Epsilon" in g for g in imp.global_issues)          # on the dashboard, no tab

def test_holidays_come_from_workingdays_and_the_location_list_does_not_blank_them(imp):
    assert imp.locations["India - Pune"] == [date(2025, 10, 2), date(2025, 10, 21), date(2026, 1, 1), date(2026, 1, 26)]
    assert len(imp.locations["US"]) == 4 and imp.locations["South Africa"] == [date(2025, 12, 16)]

def test_alpha_costco_layout(imp):
    e = next(x for x in imp.engagements if x.tab == "Alpha")
    assert e.client == "Alpha Retail" and e.business_unit == "ST21" and e.assignment_key == "SOW#12" and e.engagement_type == "time_and_materials"
    assert [a.resource for a in e.allocations] == ["Person A1", "Person A2", "Person A3"]
    a1, a2, _ = e.allocations
    assert a1.location == "US" and a2.location == "India - Pune"
    assert a1.start == date(2025, 9, 1) and a1.end == date(2027, 8, 31)          # from the SOW text
    assert [m.month for m in a1.months] == ["2025-09", "2025-10", "2025-11", "2025-12", "2026-01", "2026-02", "2026-03"]
    assert [m.kind for m in a1.months] == ["actual"] * 3 + ["forecast"] * 4
    assert a1.months[0].unit_kind == "hours" and a1.months[1].unit_kind == "days" and a1.months[1].vacation == D("1")
    assert a2.months[0].vacation == D("8") and a2.months[0].billed == D(165 * 32)

def test_alpha_flat_monthly_billing_is_imported_as_a_monthly_rate(imp, cal):
    e = next(x for x in imp.engagements if x.tab == "Alpha")
    eng, allocs, acts, notes = to_engine(e, cal)
    assert all(a.rate_basis == "monthly" and a.bill_rate in (D(173 * 120), D(173 * 32)) for a in allocs)
    assert sum("flat" in n for n in notes) == 3
    f = compute(eng, allocs, [], cal, acts)
    jan = next(m for m in f.months if m.month == "2026-01")
    assert jan.revenue == D(173 * 120 + 2 * 173 * 32)                                 # a full month = the flat amount, calendar-independent
    assert f.status == "INCOMPLETE" and any("cost per hour unknown" in m for m in f.missing)
    sep = next(m for m in f.months if m.month == "2025-09")
    assert sep.actual_revenue == D(173 * 120 + 165 * 32 + 173 * 32) and sep.actual_hours == D(173 + 165 + 173)          # hours are gross; leave hours come off

def test_beta_monthly_rates_and_allocation_pct(imp, cal):
    e = next(x for x in imp.engagements if x.tab == "Beta")
    assert e.assignment_key == "CR#01" and e.engagement_type == "msp" and e.sales_owner == "Sam"
    b1, b2 = e.allocations
    assert b1.rate_basis == "monthly" and b1.start == date(2025, 7, 1) and b1.end == date(2026, 6, 30)
    assert b2.allocation_pct == D("0.5") and b2.location == "India - Pune"
    eng, allocs, acts, notes = to_engine(e, cal)
    f = compute(eng, allocs, [], cal, acts)
    oct_ = next(m for m in f.months if m.month == "2025-10")
    assert oct_.revenue == D(5151) + D(6325) / 2                    # monthly rates × allocation, whatever the calendar says
    assert oct_.actual_revenue == D(5151) + D("3162.5")
    assert f.months[0].month == "2025-07" and f.months[-1].month == "2026-06"   # the SOW term, not the sheet's columns

def test_gamma_staffing_layout_extensions_commission_and_overtime(imp, cal):
    engs = [x for x in imp.engagements if x.tab == "Gamma Health-VMS"]
    vms = next(x for x in engs if x.assignment_key == "VMS"); po = next(x for x in engs if x.assignment_key != "VMS")
    assert vms.engagement_type == "staff_augmentation" and vms.business_unit == "IT Avalon"
    g1 = [a for a in vms.allocations if a.resource == "Person G1"]
    assert [(a.segment, a.start, a.end) for a in g1] == [(0, date(2025, 11, 3), date(2026, 2, 27)), (1, date(2026, 2, 28), date(2026, 8, 31))]
    assert g1[0].commission and g1[0].months[0].month == "2026-01"              # labels without a year → the model's year
    assert [m.kind for m in g1[0].months] == ["actual", "actual", "forecast", "forecast"]
    assert g1[0].months[0].units == D(160) and g1[0].months[0].billed == D(160 * 145)
    ot = next(a for a in vms.allocations if a.mode == "overtime")
    assert ot.location == "US"                                                      # inherited from the base row
    eng, allocs, acts, notes = to_engine(vms, cal)
    ot_alloc = next(a for a in allocs if a.person == "Person G1-OVERTIME")
    assert ot_alloc.utilization == D(0)
    assert any(x.allocation_id == ot_alloc.id and x.month == "2026-01" and x.revenue == D("870.0") for x in acts)   # 4 h × 217.5
    assert po.assignment_key == "PO 26-00389" and po.allocations[0].start == date(2026, 1, 12)

def test_closed_tab_is_flagged(imp):
    e = next(x for x in imp.engagements if x.tab == "Delta Closed")
    assert e.closed and e.allocations[0].location == "South Africa" and e.allocations[0].months[0].unit_kind == "days"

def test_dry_run_report_reads_as_a_document(imp, cal):
    md = dry_run_markdown(imp, cal)
    for s in ("## By business unit", "| ST21 |", "| IT Avalon |", "## Where the sheet forecasts past the SOW end", "## Issues by tab", "Marked CLOSED"):
        assert s in md
    rec = reconcile(imp, cal)
    assert {r["tab"] for r in rec} == {"Alpha", "Beta", "Gamma Health-VMS", "Delta Closed", "Zeta Unlisted"}
